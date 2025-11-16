"""Excel file operations for storing business listings."""
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from typing import List, Dict, Set
import os
from datetime import datetime
import config


class ExcelHandler:
    """Manages Excel file operations for business listings."""

    def __init__(self, filename: str = None):
        self.filename = filename or config.EXCEL_FILENAME
        self.existing_urls: Set[str] = set()

    def file_exists(self) -> bool:
        """Check if the Excel file already exists."""
        return os.path.exists(self.filename)

    def load_existing_urls(self) -> Set[str]:
        """Load existing listing URLs to avoid duplicates."""
        if not self.file_exists():
            return set()

        try:
            wb = load_workbook(self.filename)
            ws = wb.active

            # Find the "Listing Link" column
            link_col = None
            for idx, cell in enumerate(ws[1], 1):
                if cell.value == "Listing Link":
                    link_col = idx
                    break

            if not link_col:
                return set()

            # Extract all URLs
            urls = set()
            for row in ws.iter_rows(min_row=2, min_col=link_col, max_col=link_col):
                if row[0].value:
                    urls.add(str(row[0].value).strip())

            wb.close()
            return urls

        except Exception as e:
            print(f"Error loading existing URLs: {e}")
            return set()

    def create_new_workbook(self) -> Workbook:
        """Create a new Excel workbook with headers."""
        wb = Workbook()
        ws = wb.active
        ws.title = "Senior Care Businesses"

        # Write headers
        headers = config.EXCEL_COLUMNS
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)

            # Style the header
            cell.font = Font(bold=True, size=12)
            cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
            cell.font = Font(bold=True, size=12, color="FFFFFF")
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # Set column widths
        ws.column_dimensions['A'].width = 30  # Business Name
        ws.column_dimensions['B'].width = 50  # Financial Info
        ws.column_dimensions['C'].width = 60  # Listing Link
        ws.column_dimensions['D'].width = 40  # Broker Info
        ws.column_dimensions['E'].width = 15  # Date Added
        ws.column_dimensions['F'].width = 50  # Additional Notes

        return wb

    def append_businesses(self, businesses: List[Dict]) -> int:
        """
        Append new businesses to the Excel file, avoiding duplicates.
        Returns the number of new businesses added.
        """
        # Load existing URLs
        self.existing_urls = self.load_existing_urls()
        print(f"\nFound {len(self.existing_urls)} existing listings in Excel file")

        # Filter out duplicates
        new_businesses = [
            b for b in businesses
            if b.get('listing_link', '') not in self.existing_urls
        ]

        if not new_businesses:
            print("No new businesses to add (all are duplicates)")
            return 0

        print(f"Adding {len(new_businesses)} new businesses to Excel...")

        # Load or create workbook
        if self.file_exists():
            wb = load_workbook(self.filename)
            ws = wb.active
            next_row = ws.max_row + 1
        else:
            wb = self.create_new_workbook()
            ws = wb.active
            next_row = 2  # Start after header

        # Add new businesses
        date_added = datetime.now().strftime("%Y-%m-%d")

        for business in new_businesses:
            ws.cell(row=next_row, column=1, value=business.get('business_name', 'N/A'))
            ws.cell(row=next_row, column=2, value=business.get('financial_info', 'N/A'))
            ws.cell(row=next_row, column=3, value=business.get('listing_link', 'N/A'))
            ws.cell(row=next_row, column=4, value=business.get('broker_info', 'N/A'))
            ws.cell(row=next_row, column=5, value=date_added)
            ws.cell(row=next_row, column=6, value=business.get('additional_notes', 'N/A'))

            # Make the link clickable
            link_cell = ws.cell(row=next_row, column=3)
            link_cell.hyperlink = business.get('listing_link', '')
            link_cell.font = Font(color="0563C1", underline="single")

            next_row += 1

        # Save the workbook
        wb.save(self.filename)
        wb.close()

        print(f"✓ Successfully added {len(new_businesses)} businesses to {self.filename}")
        return len(new_businesses)

    def get_summary(self) -> Dict:
        """Get summary statistics about the Excel file."""
        if not self.file_exists():
            return {
                'total_businesses': 0,
                'file_exists': False
            }

        try:
            wb = load_workbook(self.filename)
            ws = wb.active

            total = ws.max_row - 1  # Subtract header row

            wb.close()

            return {
                'total_businesses': total,
                'file_exists': True,
                'filename': self.filename
            }

        except Exception as e:
            print(f"Error getting summary: {e}")
            return {
                'total_businesses': 0,
                'file_exists': True,
                'error': str(e)
            }
